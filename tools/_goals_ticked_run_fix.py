#!/usr/bin/env python3
"""One-shot: the `goals` sheet's single `run` Ticked cell becomes `any time`.

The Ticked column was authored with three values — `any time` (98 rows),
`game beaten` (35) and `run` (1). The one `run` row is Resourceful Rat's "Beat a
mini-game inside of a game", which is an any-time feat like the 98 rows around
it: there is nothing about beating a mini-game that waits on the outer game
being finished. Confirmed as a typo rather than a third mode, so the vocabulary
is two values and `apply_goals_sheet.py` can reject anything else outright
instead of carrying a mode nothing means.

Run once; `apply_goals_sheet.py --check` is what keeps it true afterwards.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _xlsx_surgery import Workbook  # noqa: E402

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
XLSX = os.path.join(ROOT, "tools", "Roguelikes.xlsx")

# goals!H5 — the Ticked column (H) on Resourceful Rat's row.
CELL = "H5"
WAS, NOW = "run", "any time"


def main() -> int:
    import openpyxl
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)["goals"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else "" for h in rows[0]]
    if hdr[7] != "Ticked":
        raise SystemExit("column H is %r, not Ticked — the sheet moved" % hdr[7])
    row = rows[4]
    if str(row[7]).strip() != WAS:
        print("goals!%s is already %r — nothing to do" % (CELL, row[7]))
        return 0
    print("goals!%s  %s -> %s   (%s / %s)" % (CELL, WAS, NOW, row[3], row[4]))
    with Workbook(XLSX) as wb:
        wb.replace_cells("goals", {CELL: NOW})
    return 0


if __name__ == "__main__":
    sys.exit(main())
