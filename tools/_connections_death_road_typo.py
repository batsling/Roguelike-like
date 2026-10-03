#!/usr/bin/env python3
"""One-shot: fix `Deat Road to Canada` -> `Death Road to Canada` (connections B1324).

The Nuclear Throne -> Death Road to Canada row (Kepa Auwae on Wasteland Kings,
off `docs/influence-candidates.md`) arrived with the influencee's name missing an
`h`, so `import-games-godot.py` could not resolve it and skipped the row — the
one unresolved name in that import.

Through `replace_cells` for the same reason as `_games_dq_heroes_file_typo.py`:
it edits one value and touches nothing else on the sheet.

    python3 tools/_connections_death_road_typo.py
    python3 tools/import-games-godot.py
"""

import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _xlsx_surgery import Workbook  # noqa: E402

XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Roguelikes.xlsx")

SHEET = "connections"
INFLUENCER_CELL = "A1324"
INFLUENCER = "Nuclear Throne"
CELL = "B1324"
WAS = "Deat Road to Canada"
NOW = "Death Road to Canada"


def main():
    # The ref is only trustworthy if the row is still the row this was written
    # against; a sort of the sheet would move it.
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)[SHEET]
    influencer = str(ws[INFLUENCER_CELL].value or "").strip()
    current = str(ws[CELL].value or "").strip()
    if influencer != INFLUENCER:
        raise SystemExit("%s reads %r, not %r — the sheet has been re-sorted; find "
                         "the row again before rerunning" % (INFLUENCER_CELL, influencer, INFLUENCER))
    if current == NOW:
        raise SystemExit("%s already reads %r — nothing to do" % (CELL, NOW))
    if current != WAS:
        raise SystemExit("%s reads %r, not the cell this edit was written against "
                         "— re-read the row before rerunning" % (CELL, current))

    with Workbook(XLSX) as wb:
        wb.replace_cells(SHEET, {CELL: NOW})
    print("%s!%s: %s -> %s" % (SHEET, CELL, WAS, NOW))


if __name__ == "__main__":
    main()
