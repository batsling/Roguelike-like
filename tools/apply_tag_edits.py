#!/usr/bin/env python3
"""Write the tags edited IN GAME into the `games` sheet's Tags column.

The game lets the player add tags to a game and take them off (the Collection's
game page and the run map's game card; `scripts/autoload/GameTags.gd`), and its
Export button writes what the sheet doesn't have yet to `tools/tag_edits.json`
(or the user folder, from an exported build). This applies that file:

    python3 tools/apply_tag_edits.py                  # tools/tag_edits.json -> sheet, then re-import
    python3 tools/apply_tag_edits.py path/to/file.json
    python3 tools/apply_tag_edits.py --dry-run        # say what would change, write nothing
    python3 tools/apply_tag_edits.py --no-import      # write the sheet, skip import-games-godot.py

For each game in the file: the sheet's tags, minus what it says to remove, plus
what it says to add (appended, in the order added), written back as one
comma-separated cell. It is IDEMPOTENT: an addition the cell already has and a
removal it already lacks change nothing, so applying a file twice, or applying
an old one, is harmless. Games are matched by name (the sheet's key), falling
back to the id slug the importer derives, so a renamed game is reported rather
than guessed at.

Through `_xlsx_surgery.set_cells`, never openpyxl's save: a round trip through
openpyxl drops the workbook's charts. Only column F of the matched rows is
written; every other cell, the formula columns included, is copied through.

Then it runs `import-games-godot.py`, which bakes the new tags into
`data/games/` — and once the game loads that, its own pending list drops every
edit the sheet now has (`GameTags._prune`), so nobody clears anything by hand.
Delete the JSON in the same commit; leaving it is harmless, but it is spent.
"""

import json
import os
import re
import subprocess
import sys

import openpyxl  # read-only here; the write goes through _xlsx_surgery

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from _xlsx_surgery import Workbook  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "tools", "Roguelikes.xlsx")
DEFAULT = os.path.join(ROOT, "tools", "tag_edits.json")
SHEET = "games"
NAME_COL, TAGS_COL = "A", "F"   # Name | Year | Type | Connected? | Influencer? | Tags
HEADERS = ["Name", "Year", "Type", "Connected?", "Influencer?", "Tags"]


def normalize(tag) -> str:
    """The spelling GameTags.normalize keeps: lower case, single spaces, no commas."""
    return re.sub(r"\s+", " ", str(tag or "").replace(",", " ")).strip().lower()


def slug(name: str) -> str:
    """import-games-godot.py's id_for: what `data/games/<id>.tres` is called."""
    name = re.sub(r"[\[\]]", "", name)
    return "_".join(p.lower() for p in re.sub(r"[^a-zA-Z0-9\s]", " ", name).split() if p)


def split_tags(cell) -> list:
    out = []
    for t in str(cell or "").split(","):
        n = normalize(t)
        if n and n not in out:
            out.append(n)
    return out


def sheet_rows():
    """[(row number, name, tags cell)] for every named row of the games sheet."""
    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[SHEET]
    rows, header = [], None
    for cells in ws.iter_rows():
        real = [c for c in cells if hasattr(c, "row")]
        if not real:
            continue
        r = real[0].row
        values = {c.column_letter: c.value for c in real if hasattr(c, "column_letter")}
        if header is None:
            header = [str(values.get(chr(65 + i)) or "").strip() for i in range(len(HEADERS))]
            if header != HEADERS:
                sys.exit("games headers are %r, expected %r — the sheet's shape moved; "
                         "update this script" % (header, HEADERS))
            continue
        name = str(values.get(NAME_COL) or "").strip()
        if name:
            rows.append((r, name, values.get(TAGS_COL)))
    return rows


def plan(edits: dict, rows: list):
    """({cell ref: new value}, [report lines], [unmatched game names])."""
    by_name = {name: (r, cell) for r, name, cell in rows}
    by_slug = {slug(name): (r, cell, name) for r, name, cell in rows}
    cells, report, missing = {}, [], []
    for g in edits.get("games", []):
        name = str(g.get("name") or "").strip()
        hit = by_name.get(name)
        sheet_name = name
        if hit is None:
            s = by_slug.get(str(g.get("id") or "")) or by_slug.get(slug(name))
            if s is None:
                missing.append(name or str(g.get("id")))
                continue
            hit, sheet_name = (s[0], s[1]), s[2]
        r, cell = hit
        before = split_tags(cell)
        drop = {normalize(t) for t in g.get("remove", [])}
        after = [t for t in before if t not in drop]
        for t in g.get("add", []):
            n = normalize(t)
            if n and n not in after:
                after.append(n)
        if after != before:
            cells["%s%d" % (TAGS_COL, r)] = ", ".join(after)
            report.append("  %-40s %s  ->  %s" % (sheet_name[:40], ", ".join(before) or "(none)",
                                                    ", ".join(after) or "(none)"))
    return cells, report, missing


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    path = args[0] if args else DEFAULT
    if not os.path.exists(path):
        sys.exit("no tag edits at %s — export them from the game first (the Export button "
                 "under any game's tags)" % os.path.relpath(path, ROOT))
    edits = json.load(open(path, encoding="utf8"))
    cells, report, missing = plan(edits, sheet_rows())
    print("%s: %d game(s) in the file, %d cell(s) to change" % (
        os.path.relpath(path, ROOT), len(edits.get("games", [])), len(cells)))
    for line in report:
        print(line)
    for name in missing:
        print("  NOT ON THE SHEET (renamed or removed?): %s" % name)
    if dry or not cells:
        return 1 if missing else 0
    with Workbook(XLSX) as wb:
        wb.set_cells(SHEET, cells)
    print("wrote %d cell(s) to the games sheet" % len(cells))
    if "--no-import" not in sys.argv:
        rc = subprocess.call([sys.executable, os.path.join(ROOT, "tools", "import-games-godot.py")])
        if rc != 0:
            return rc
    return 1 if missing else 0


if __name__ == "__main__":
    sys.exit(main())
