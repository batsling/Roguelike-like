#!/usr/bin/env python3
"""Research every game on the chart for what it could give this one.

ONE ENTRY POINT for every kind of research the project does. Each kind asks one
question of a real game:

    connections  which chart games influenced it, in the developer's own words
    tags         which of the sheet's themes it carries (tools/tag_research.py)
    goals        which of its enemies and bosses would make a good goal-enemy
    loot         which of its items could become loot here, and of which kind
    events       which of its events or encounters could become an event
    characters   which of its playable characters could become a character
    statuses     which of its status effects or curses could become one here
    locations    which of its places or map objects could become a location
                 or an object

The full guide, with the rules each kind's rows have to pass, is
`docs/research.md`. In short:

  * Candidates live in `research/<kind>.csv`, one per kind, CHECKED IN. They
    are the source of truth, readable in a diff.
  * `tools/Research.xlsx` is where the owner reviews them: a sheet per kind, with
    an `Owner` column (yes / no) and `Owner Notes`. `build` writes the sheets from
    the CSVs; `sync` reads the owner's edits back. Nothing here ever writes
    `tools/Roguelikes.xlsx`: a candidate becomes content only once the owner
    has ticked it.
  * `research/ledger.json` records which games have been researched for which
    kind, so no session repeats another's work and `status` can show the gaps.
  * `research/wikis.json` records each game's wiki, found once by `wikis` and
    open to correction by hand. The wikis are where goals, loot, events,
    characters, statuses and locations come from.

Commands:

    python3 tools/research.py status              # coverage per kind, and what is next
    python3 tools/research.py next goals -n 5     # the next games to research for a kind
    python3 tools/research.py wikis               # find wikis for games that have none recorded
    python3 tools/research.py brief "Hades"       # everything about a game, for writing its candidates
    python3 tools/research.py page "Hades" --category Boons --chars 300
    python3 tools/research.py page "Hades" "Charon" "Hypnos"
    python3 tools/research.py mark goals "Hades" "Brotato" --note "12 rows"
    python3 tools/research.py check               # the CSVs, the ledger and wikis.json are sound
    python3 tools/research.py build               # CSVs -> Research.xlsx
    python3 tools/research.py sync                # owner's edits in Research.xlsx -> CSVs, then build
    python3 tools/research.py new                 # every kind, for the games just added

Network work is cached in `.research_work/` (gitignored).
"""

import argparse
import concurrent.futures as cf
import csv
import datetime
import html
import html.parser
import json
import os
import re
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.parse
import urllib.request

import openpyxl  # Roguelikes.xlsx is opened READ-ONLY here; only Research.xlsx is written

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import research_book as rb  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TOOLS = os.path.join(ROOT, "tools")
STORE = os.path.join(ROOT, "research")
LEDGER = os.path.join(STORE, "ledger.json")
WIKIS = os.path.join(STORE, "wikis.json")
SHEET = os.path.join(TOOLS, "Roguelikes.xlsx")
WORK = os.path.join(ROOT, ".research_work")
UA = {"User-Agent": "Roguelike-like research (https://github.com/batsling/roguelike-like)"}
TODAY = datetime.date.today().isoformat()

# ── the kinds ───────────────────────────────────────────────────────────────

# The columns every candidate kind ends with. The sheet is a BRAINSTORM the owner
# copies from by hand, so after the target sheet's own columns come three that
# make the idea readable at a glance: `What it is there` (the thing in its own
# game), `Builds on` (the system of THIS game it plugs into: "Fire tile", "coin
# trinkets", "add_goal event") and `Needs` (blank if it rides what exists,
# otherwise the new mechanic it would take). Then `Why it fits`, `Source` (the
# page), `Confidence` (`ok`, or `check`: written from memory or a summary, or a
# detail nobody could confirm), `Status` (STATUSES), the owner's two columns,
# and a stable `ID`, which is how `sync` matches a workbook row to its CSV line
# however the owner sorts or filters the sheet.
#
# `Clones` (after `Needs`) holds the ideas this row beat: when two candidates,
# from any games, overlap in mechanic or goal, the more interesting one is the
# row and the other is written here as `Name (Game): how it differs`, entries
# split by `; `, instead of being a second row or being dropped. It came after
# the rest, so `sync` still accepts a workbook built before it (LATE_COLUMNS).
IDEA = ["What it is there", "Builds on", "Needs", "Clones"]
STAGE = ["What it is there", "Why it fits", "Builds on", "Needs", "Clones", "Source", "Confidence", "Status",
         "Owner", "Owner Notes", "ID"]
LATE_COLUMNS = {"Clones"}

EVENT_CHOICES = [f"{c} {i}" for i in range(1, 7) for c in ("Choice", "Repeat", "Result", "Effect")]

KINDS = {
    "connections": dict(
        title="Connections: which chart games influenced which",
        columns=["Influencer", "Influencee", "Game", "Dev/Series", "Strength", "Text", "Source", "Heading",
                 "Status", "Owner", "Owner Notes", "ID"],
        name="Influencee", targets=["connections"]),
    "goals": dict(
        title="Goal-enemies and bosses: a game's monsters, each carrying a goal",
        # The `enemies` / `bosses` sheets' own columns in their own order, then
        # the staging ones; tools/check_goal_candidates.py audits this file.
        columns=["Sheet", "Name", "Type", "Difficulty", "Size", "Game", "Health", "Damage", "Goal Type",
                 "Goal", "Ability", "File", "Tag", "Phases", "Ticked", "Count", "Confidence",
                 "Why this pairing"] + IDEA + ["Source", "Status", "Owner", "Owner Notes", "ID"],
        name="Name", sheets=["enemies", "bosses"], categories="goals"),
    "loot": dict(
        title="Loot: items, trinkets, cards, weapons, evolutions, bags, wands, potions and scrolls",
        # One file for every loot sheet: the columns they share, `Effect` left
        # for the owner, and `Extra` for the few a single sheet has (a wand's
        # Charges, a weapon's Aim / Area / Goal, an evolution's Requirements) as
        # `key=value; key=value`. `check` holds weapons and evolutions to theirs.
        columns=["Sheet", "Name", "Game", "Rarity", "Type", "Size", "Description", "Effect", "Tags", "File",
                 "Extra"] + STAGE,
        name="Name", sheets=["items", "trinkets", "cards", "weapons", "evolutions", "bags", "wands", "potions",
                                     "scrolls"],
        categories="loot"),
    "events": dict(
        title="Events: a prompt between games, with choices",
        columns=["Event", "Game", "Tier", "Where", "Requirement", "Trigger", "Rarity", "Image", "Prompt",
                 "Opens With", "Goal Met", "Goal Missed", "Chance Won", "Chance Lost"] + EVENT_CHOICES + STAGE,
        name="Event", targets=["events"], categories="events"),
    "characters": dict(
        title="Characters: playable heroes, with a level-up goal",
        columns=["Name", "Game", "Health", "Gold", "Bash", "Dash", "Push", "Transmute", "Scramble", "Bombs",
                 "Keys", "Random", "Level Up", "Reward", "Description", "Starting loadout", "File"] + STAGE,
        name="Name", targets=["characters"], categories="characters"),
    "statuses": dict(
        title="Statuses and curses",
        columns=["Sheet", "Name", "Game", "Type", "On Player", "On Enemy", "Combat",
                 "Condition", "Penalty", "Timer", "Image"] + STAGE,
        name="Name", sheets=["statuses", "curses"], categories="statuses"),
    "locations": dict(
        title="Locations and objects: places on the map, and things standing in them",
        columns=["Sheet", "Name", "Game", "Difficulty", "Goal Type", "Goal", "Goal Effect", "Tag", "Rarity",
                 "Description", "Choices", "Image"] + STAGE,
        # `shopkeepers` has no sheet yet: ShopPanel2 reads a `shopkeeper` field
        # and nothing is authored for it (spec §14.5), so ideas wait here.
        # `tiles` and `units` are the battlefield's terrain and placed things
        # (Fire, Web, Landmine): what stands ON the board rather than on the map.
        name="Name", sheets=["locations", "objects", "shopkeepers", "tiles", "units"], categories="locations"),
}
# Kinds whose rows are written by hand from a game's sources, game by game.
# `connections` is researched by tools/influence_research.py, and `tags` by
# tools/tag_research.py, which writes its own sheets and needs no CSV.
WIKI_KINDS = [k for k in KINDS if KINDS[k].get("categories")]
LEDGER_KINDS = list(KINDS)

STATUSES = {
    "to review": "waiting for the owner",
    "waiting for game row": "names a game not on the `games` sheet yet",
    "lead": "worth following up; not a candidate yet",
    "source check": "a row already on the sheet whose source needs a look",
    "not an influence": "checked, and the answer is no (kept so nobody adds it later)",
    "nothing found": "searched, nothing usable (kept so it is not searched again)",
    "on sheet": "the row is in Roguelikes.xlsx now",
}
OWNER = {"", "yes", "no"}
CONFIDENCE = {"", "ok", "check", "?"}  # `?` is the goal file's older spelling of `check`

# Where each target sheet keeps its name and its game. `scrolls` has two
# `Game` columns; the first is the scroll's.
TARGET_COLS = {
    "enemies": ("Name", "Game"), "bosses": ("Name", "Game"), "items": ("Name", "Reference"),
    "trinkets": ("Name", "Game"), "cards": ("Name", "Game"), "weapons": ("Name", "Game"),
    "bags": ("Name", "Game"), "wands": ("Name", "Game"), "potions": ("Name", "Reference"),
    "scrolls": ("Scrolls", "Game"), "events": ("Event", "Game"), "characters": ("Name", "Game"),
    "statuses": ("Name", "Game"), "curses": ("Curse", "Game"), "locations": ("Name", "Game"),
    "objects": ("Name", "Game"), "tiles": ("Name", "Game"), "units": ("Name", "Game"),
    # No game column: an evolution is named for the weapon it becomes.
    "evolutions": ("Name", "Game"),
}

# A weapon or evolution candidate carries its sheet's own columns in `Extra`
# (docs/loot-passives.md §12-13). These are the keys `check` insists on.
EXTRA_KEYS = {
    "weapons": ("Aim", "Area", "Type", "Goal", "Charge"),
    "evolutions": ("Requirement 1", "Requirement 2", "Outcome"),
}
WEAPON_AIMS = {"any", "front", "back", "enemy", "none", "random"}
EVOLUTION_OUTCOMES = {"Consume All", "Consume None"}

# Weapons and evolutions are asked of every game (the owner's standing
# request), so `brief` looks for them by name on the live wiki rather than
# trusting the general loot categories to surface them.
WEAPON_CATS = re.compile(r"weapons?|\bguns?\b|firearms?|swords?|\bmelee\b|\branged\b|\bbows?\b|"
                         r"staff|staves|blades?|spears?|\baxes?\b|daggers?|hammers?|launchers?|"
                         r"\borbs?\b|\bballs?\b", re.I)  # Peglin's orbs and Ball x Pit's balls are their weapons
EVOLUTION_CATS = re.compile(r"evolutions?|evolved|synerg|\bfusions?\b|transformations?|recipes?|combos?|"
                            r"combinations?|crafting|upgrade paths?", re.I)
WEAPON_PAGES = ["Weapons", "Guns", "Melee Weapons", "Ranged Weapons", "Weapon", "Swords", "Orbs", "Balls"]
EVOLUTION_PAGES = ["Evolutions", "Evolution", "Weapon Evolutions", "Synergies", "Synergy", "Transformations",
                   "Fusion", "Fusions", "Recipes", "Crafting", "Combinations", "Combos"]

# What a wiki category is about, by its name. Used for two things: ranking games
# by how much their wiki documents (`next`), and listing what to read (`brief`).
CATEGORY = {
    "goals": r"enem(y|ies)|monsters?|bosses|\bboss\b|creatures?|\bmobs?\b|minions|foes|bestiary|villains",
    "loot": r"\bitems?\b|relics?|\bcards?\b|weapons?|\bguns?\b|potions?|scrolls?|wands?|\brings?\b|amulets?|"
            r"trinkets?|evolutions?|synerg|fusions?|transformations?|recipes?|"
            r"artifacts?|artefacts?|equipment|consumables?|passives?|actives?|pickups?|jokers?|upgrades|perks?|"
            r"blessings?|boons?|\bgear\b|armou?r|charms?|curios?|\bfood\b|treasures?|tools\b|spells?|augments?",
    "events": r"\bevents?\b|encounters?|\bquests?\b|dialogues?|\bchoices\b",
    "characters": r"characters?|\bclasses\b|heroes|playable|survivors|protagonists?|\braces\b|companions",
    "statuses": r"status|debuffs?|\bbuffs?\b|curses?|conditions|afflictions?|ailments?|mutations?|keywords?|"
                r"intrinsics?|effects\b",
    "locations": r"locations?|\brooms?\b|\bareas?\b|biomes?|\bfloors?\b|\bstages?\b|\bzones?\b|regions?|"
                 r"dungeons?|shrines?|altars?|\bshops?\b|structures?|\bobjects?\b|machines?|landmarks?|"
                 r"\bmaps?\b|chapters?|\bworlds?\b|terrain|features",
}
NOT_CONTENT = re.compile(
    r"images?|icons?|sprites?|templates?|\bfiles?\b|stubs?|navbox|infobox|achievements?|music|soundtrack|"
    r"screenshots?|videos?|disambiguation|deletion|\busers?\b|maintenance|pages with|articles|\bcss\b|"
    r"\bjs\b|cleanup|redirects?|candidates|translat|\bstub\b|hidden|tracking|lua|modules?|gifs?|"
    r"animations?|portraits?|artwork|concept art|patch|updates? \d|version|cut content|unused|"
    r"navigation|modding|\bdata\b|\btables?\b|population|lists? of|defunct|missing|"
    r"\b(ja|zh|ko|ru|de|fr|es|pt|pl|it)\b", re.I)

# The list articles a wiki keeps when a kind has no category of its own — a
# wiki's status effects are nearly always one page ("Status Effects", "Buffs"),
# not a category. `wikis` records which of these exist; `brief` lists them.
LIST_PAGES = {
    "goals": ["Enemies", "Monsters", "Bosses", "Bestiary", "Creatures", "Mobs", "Minibosses"],
    "loot": ["Items", "Relics", "Weapons", "Cards", "Potions", "Scrolls", "Wands", "Trinkets", "Jokers",
             "Artifacts", "Equipment", "Consumables", "Upgrades", "Boons", "Pickups", "Rings", "Amulets",
             "Spells", "Tools", "Food", "Charms", "Perks", "Blessings", "Curios", "Augments", "Guns",
             "Evolutions", "Synergies", "Transformations", "Recipes"],
    "events": ["Events", "Encounters", "Random events", "Random Events", "Quests", "Shrines", "NPCs"],
    "characters": ["Characters", "Classes", "Heroes", "Playable characters", "Survivors", "Races", "Roles"],
    "statuses": ["Status Effects", "Status effects", "Status Effect", "Statuses", "Status", "Buffs", "Debuffs",
                 "Buffs and Debuffs", "Curses", "Conditions", "Powers", "Afflictions", "Ailments", "Mutations",
                 "Intrinsics", "Keywords"],
    "locations": ["Locations", "Rooms", "Areas", "Biomes", "Floors", "Stages", "Chapters", "Levels", "Zones",
                  "Shrines", "Altars", "Shops", "Machines", "Objects", "Map", "Dungeon", "Structures"],
}

# Wikis the address-guessing can't find: the community runs them on its own
# domain. A game's entry in research/wikis.json marked `"set": "by hand"`
# overrides discovery the same way; this list is just the ones known up front.
KNOWN_WIKIS = {
    "NetHack": "https://nethackwiki.com/api.php",
    "Dungeon Crawl Stone Soup": "http://crawl.chaosforge.org/api.php",
    "Dwarf Fortress": "https://dwarffortresswiki.org/api.php",
    "Caves of Qud": "https://wiki.cavesofqud.com/api.php",
}


# ── files ───────────────────────────────────────────────────────────────────

def csv_path(kind):
    return os.path.join(STORE, f"{kind}.csv")


def read_rows(kind):
    path = csv_path(kind)
    if not os.path.exists(path):
        return []
    with open(path, newline="", encoding="utf-8") as fh:
        rd = csv.DictReader(fh)
        return [{c: (r.get(c) or "").strip() for c in rd.fieldnames} for r in rd]


def write_rows(kind, rows):
    cols = KINDS[kind]["columns"]
    os.makedirs(STORE, exist_ok=True)
    with open(csv_path(kind), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=cols, lineterminator="\n")
        w.writeheader()
        for r in rows:
            w.writerow({c: r.get(c, "") for c in cols})


def load_json(path, default):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf8") as fh:
        return json.load(fh)


def save_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf8") as fh:
        json.dump(data, fh, indent=1, sort_keys=True, ensure_ascii=False)
        fh.write("\n")


def load_ledger():
    """{game: {kind: "YYYY-MM-DD[: note]"}}."""
    return load_json(LEDGER, {})


# A ledger entry starting `redo:` is research that has to be done again: on the
# record, but not done. The 80 games of the old goal passes are the case it was
# made for: written from search summaries while the wikis were unreachable, so a
# pass with the wiki open checks each row, fixes it, and adds what was missed.
REDO = "redo:"


def researched(ledger, game, kind):
    v = ledger.get(game, {}).get(kind)
    return v is not None and not v.startswith(REDO)


def mark(games, kind, note=""):
    ledger = load_ledger()
    for g in games:
        ledger.setdefault(g, {})[kind] = TODAY + (f": {note}" if note else "")
    save_json(LEDGER, ledger)


def slug(s):
    s = s.lower().replace("&", "and")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-") or "x"


# ── the catalog ─────────────────────────────────────────────────────────────

_SHEET = None


def workbook():
    global _SHEET
    if _SHEET is None:
        _SHEET = openpyxl.load_workbook(SHEET, read_only=True, data_only=True)
    return _SHEET


def sheet_table(name):
    """[row dict] of a Roguelikes.xlsx sheet; a repeated header keeps its first column."""
    rows = list(workbook()[name].iter_rows(values_only=True))
    head = [rb.norm(h) for h in rows[0]]
    first = {}
    for i, h in enumerate(head):
        first.setdefault(h, i)
    return [{h: rb.norm(r[i]) if i < len(r) else "" for h, i in first.items() if h}
            for r in rows[1:] if r and r[0] is not None]


def catalog():
    """{name: {year, type, owned, steam}} in sheet order."""
    out = {}
    for r in sheet_table("games"):
        out[r["Name"]] = {"year": r.get("Year", ""), "type": r.get("Type", ""),
                          "owned": r.get("Owned", "").lower() == "yes", "steam": r.get("Steam Page", "")}
    return out


def on_sheet_index():
    """{target sheet: {lower name: game}} for every sheet a candidate can land on."""
    out = {}
    for sheet, (ncol, gcol) in TARGET_COLS.items():
        out[sheet] = {r.get(ncol, "").lower(): r.get(gcol, "") for r in sheet_table(sheet) if r.get(ncol)}
    return out


def connection_pairs():
    return {(r["Influencer"].lower(), r["Influencee"].lower()) for r in sheet_table("connections")}


def target_sheet(kind, row):
    k = KINDS[kind]
    if "sheets" in k:
        return row.get("Sheet", "")
    return k["targets"][0]


# ── HTTP and wikis ──────────────────────────────────────────────────────────

# Politeness. A first run of `wikis` fired 8 threads at Fandom unpaced, and
# Fandom answered by rate-limiting this client for a while. So every request
# waits its turn per host family (all of *.fandom.com counts as one), and a 429
# is answered by sleeping as long as the server asks.
PACE = 0.5  # seconds between requests to one host family
_last, _pace_lock = {}, threading.Lock()


def _host_family(url):
    host = urllib.parse.urlparse(url).netloc
    return ".".join(host.split(".")[-2:])


def _wait_turn(url):
    fam = _host_family(url)
    with _pace_lock:
        now = time.monotonic()
        at = max(now, _last.get(fam, 0) + PACE)
        _last[fam] = at
    if at > now:
        time.sleep(at - now)


def fetch(url, timeout=30, tries=4):
    for i in range(tries):
        _wait_turn(url)
        try:
            req = urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.status, r.read().decode("utf8", "replace")
        except urllib.error.HTTPError as e:
            if e.code in (429, 502, 503) and i < tries - 1:
                after = e.headers.get("Retry-After", "") if e.headers else ""
                time.sleep(min(int(after), 120) if after.isdigit() else 2 ** (i + 2))
                continue
            return e.code, ""
        except Exception:
            if i < tries - 1:
                time.sleep(2 ** i)
                continue
            return 0, ""
    return 0, ""


def api(base, **params):
    """One MediaWiki API call. wiki.gg answers a busy client with HTTP 200 and
    `{"error": {"code": "ratelimited"}}`, which used to read as "no page"; it
    is waited out (up to a few minutes) instead, and reported if it persists."""
    params.setdefault("format", "json")
    for wait in (15, 45, 90, 0):
        status, body = fetch(base + "?" + urllib.parse.urlencode(params))
        if status != 200:
            return None
        try:
            d = json.loads(body)
        except ValueError:
            return None
        if (d.get("error") or {}).get("code") != "ratelimited":
            return d
        if not wait:
            print(f"  rate-limited by {urllib.parse.urlparse(base).netloc}; try again later", file=sys.stderr)
            return None
        time.sleep(wait)
    return None


def _title_norm(s):
    s = s.lower().replace("&", "and")
    s = re.sub(r"\b(wiki|wikia|the|official|community|gamepedia|fandom|en)\b", " ", s)
    s = re.sub(r"[^a-z0-9]", "", s)
    # "NetHackWiki", "WikiPowder", "Downwell Wikia": the word run into the name.
    return re.sub(r"^wiki|wikia?$", "", s)


SEQUEL = re.compile(r"\s+(\d+|ii|iii|iv|v|vi|remastered|returns|classic|deluxe|hd|plus|\+|redux|"
                    r"reborn|rebirth|legacy|zero)$", re.I)


def wiki_guesses(game):
    """api.php addresses the game's wiki is likely to have: its full title, the
    part before a subtitle, and the series name a sequel shares a wiki under."""
    main = re.split(r":| - | – |—", game)[0].strip()
    titles = [game, main, SEQUEL.sub("", main)]
    subs = []
    for t in titles:
        t = t.lower().replace("&", "and")
        for v in (t, re.sub(r"^the\s+", "", t)):
            for s in (re.sub(r"[^a-z0-9]", "", v), re.sub(r"[^a-z0-9]+", "-", v).strip("-")):
                if s and s not in subs:
                    subs.append(s)
    out = []
    for s in subs:
        out += [f"https://{s}.fandom.com/api.php", f"https://{s}.wiki.gg/api.php"]
    return out


def siteinfo(base):
    d = api(base, action="query", meta="siteinfo", siprop="general|statistics")
    if not d or "query" not in d:
        return None
    g, st = d["query"]["general"], d["query"].get("statistics", {})
    return {"sitename": g.get("sitename", ""), "articles": st.get("articles", 0),
            "url": g.get("base", "").rsplit("/wiki/", 1)[0] or g.get("server", "")}


def match_quality(game, sitename):
    """`sure` when the wiki is named after the game — its full title, subtitle
    and all — and `check` when one name contains the other: a series wiki, or a
    wiki for the game a subtitle hangs off ("Dragon Quest Wiki" for Dragon
    Quest: Young Yangus covers forty games). None otherwise."""
    a = _title_norm(re.split(r":| - | – ", game)[0])
    full = _title_norm(game)
    b = _title_norm(sitename)
    if not b:
        return None
    if b == full:
        return "sure"
    if b == a:
        return "check"
    if len(min(a, b, key=len)) >= 4 and (a in b or b in a or full in b):
        return "check"
    return None


def main_page(base):
    """The wiki's main page as raw wikitext, or "". Raw because most main pages
    are built out of templates and the words are in their arguments; and not
    `action=parse`, which Fandom's Cloudflare answers with a challenge page."""
    d = api(base, action="query", meta="siteinfo")
    if not d:
        return ""
    title = d["query"]["general"].get("mainpage", "Main Page")
    q = api(base, action="query", prop="revisions", rvprop="content", rvslots="main", titles=title,
            redirects="1") or {}
    for p in q.get("query", {}).get("pages", {}).values():
        revs = p.get("revisions") or [{}]
        slot = revs[0].get("slots", {}).get("main", revs[0])
        return slot.get("*") or slot.get("content") or ""
    return ""


def blurb_of(raw):
    """A line a person can judge the wiki by."""
    return " ".join(re.sub(r"\[facts\].*", "", wikitext_to_text(raw)).split())[:240]


def signal_of(year, raw):
    """What the main page says, kept so a later run can re-judge without
    fetching: "rogue" if it says roguelike/-lite, "year" if it names the
    release year, "" if neither, None if it could not be read."""
    if not raw:
        return None
    if "rogue" in raw.lower():
        return "rogue"
    return "year" if year and str(year) in raw else ""


def looks_like_the_game(year, raw, counts, signal=False):
    """Is a wiki named like the game about it? Guessing an address finds
    whatever owns the name: "Ragnarok" found a Ragnarök fan wiki, "Ringer" a
    TV show's, "Omega" a web series'. A roguelike's wiki nearly always says
    roguelike (or -lite) or the release year on its main page, or at least
    keeps pages of enemies or items. ("game" is no signal: a TV wiki's
    navigation says it.) What fails all three is `check`, not rejected; a
    main page that could not be read is no evidence either way."""
    sig = signal_of(year, raw) if signal is False else signal
    if sig is None:
        return True
    content = (counts or {}).get("goals", 0) + (counts or {}).get("loot", 0)
    return bool(sig) or content >= 5


def count_categories(base, cap_pages=20):
    """{kind: [(category, pages)]} for the categories that look like content."""
    found = {k: [] for k in CATEGORY}
    params = dict(action="query", list="allcategories", acprop="size", aclimit="500", acmin="2")
    for _ in range(cap_pages):
        d = api(base, **params)
        if not d:
            break
        for c in d.get("query", {}).get("allcategories", []):
            name, pages = c.get("*", ""), c.get("pages", 0)
            if pages < 2 or NOT_CONTENT.search(name):
                continue
            for kind, pat in CATEGORY.items():
                if re.search(pat, name, re.I):
                    found[kind].append((name, pages))
        cont = d.get("continue", {}).get("accontinue")
        if not cont:
            break
        params["accontinue"] = cont
    return {k: sorted(v, key=lambda x: -x[1]) for k, v in found.items()}


def list_pages(base):
    """{kind: [titles]} of LIST_PAGES that exist on this wiki."""
    want = sorted({t for ts in LIST_PAGES.values() for t in ts})
    have = set()
    for i in range(0, len(want), 50):
        d = api(base, action="query", titles="|".join(want[i:i + 50]), redirects="1")
        if not d:
            continue
        q = d.get("query", {})
        back = {r["to"]: r["from"] for r in q.get("redirects", [])}
        for p in q.get("pages", {}).values():
            if "missing" not in p and "invalid" not in p:
                have.add(back.get(p["title"], p["title"]))
    return {k: [t for t in ts if t in have] for k, ts in LIST_PAGES.items() if any(t in have for t in ts)}


def discover(game, year=""):
    """The best wiki for a game, or a record saying none was found."""
    cands = []
    if game in KNOWN_WIKIS:
        info = siteinfo(KNOWN_WIKIS[game])
        if info:
            cands.append((KNOWN_WIKIS[game], info, "known"))
    for base in wiki_guesses(game):
        info = siteinfo(base)
        if info:
            q = match_quality(game, info["sitename"])
            if q:
                cands.append((base, info, q))
    if not cands:
        return {"wiki": None, "checked": TODAY}
    # A named-after-the-game wiki beats a series one; then the bigger, which is
    # the live one when a community has moved (most moved Fandom -> wiki.gg).
    cands.sort(key=lambda c: (c[2] not in ("known", "sure"), -c[1]["articles"]))
    base, info, q = cands[0]
    raw = main_page(base)
    cats = count_categories(base)
    counts = {k: sum(p for _, p in v) for k, v in cats.items()}
    if q == "known":
        q = "sure"
    elif q == "sure" and not looks_like_the_game(year, raw, counts):
        q = "check"
    rec = {"wiki": info["url"], "api": base, "sitename": info["sitename"], "articles": info["articles"],
           "match": q, "checked": TODAY, "blurb": blurb_of(raw), "signal": signal_of(year, raw),
           "counts": counts,
           "categories": {k: [n for n, _ in v[:12]] for k, v in cats.items() if v},
           "pages": list_pages(base)}
    others = [c[0] for c in cands[1:] if c[0] != base]
    if others:
        rec["also"] = others
    return rec


SCORE_CAP = 300


def wiki_score(rec, kind=None):
    """How much a game's wiki documents, for ranking: pages in the kind's
    categories, plus a little for each list article (where a kind with no
    category, statuses especially, actually lives)."""
    if not rec or not rec.get("wiki"):
        return 0
    counts, pages = rec.get("counts", {}), rec.get("pages", {})
    kinds = [kind] if kind else list(CATEGORY)
    # Capped: past a few hundred pages a wiki is "well documented" and more is
    # mostly a big franchise (a series wiki counts every game in it). A series
    # match is worth half, since most of what it documents is other games.
    score = sum(min(counts.get(k, 0), SCORE_CAP) + 25 * len(pages.get(k, [])) for k in kinds)
    return score // 2 if rec.get("match") == "check" else score


# ── reading a wiki page as text ─────────────────────────────────────────────

# Infobox fields that are a wiki's machinery rather than facts about the thing.
BOX_SKIP = re.compile(r"image|icon|sprite|file|caption|gallery|^id$|^nav|colou?r|render|^tile|bits|weight|"
                      r"version|categor|inherit|supported|unidentified|complexity|canbuild|candisassemble|"
                      r"isthrown|^type$|feature|^title$|^name$|footnote|noindent|fullwidth|centered|border", re.I)


def _strip_templates(text, box=None):
    """Drop {{templates}}. An infobox's `key = value` fields go into `box`
    (that is where a wiki keeps a monster's health or an item's effect), and a
    template that just wraps a name ({{favilink|Salve injector}}) keeps it."""
    out, i, n = [], 0, len(text)
    while i < n:
        if text.startswith("{{", i):
            depth, j = 0, i
            while j < n:
                if text.startswith("{{", j):
                    depth += 1
                    j += 2
                elif text.startswith("}}", j):
                    depth -= 1
                    j += 2
                    if depth == 0:
                        break
                else:
                    j += 1
            body = text[i + 2:j - 2]
            parts = body.split("|")
            if len(parts) > 3 and sum("=" in p for p in parts) >= 3:
                for p in parts[1:]:
                    if "=" in p and box is not None:
                        k, v = p.split("=", 1)
                        k, v = k.strip(), " ".join(_strip_templates(v).split())
                        if v and not BOX_SKIP.search(k) and len(v) < 300 and "}}" not in v:
                            box.append(f"{k}: {v}")
            elif len(parts) > 1 and "=" not in parts[1] and len(parts[1]) < 120:
                out.append(_strip_templates(parts[1]).strip())
            i = j
        else:
            out.append(text[i])
            i += 1
    return "".join(out)


def wikitext_to_text(t):
    """A page as prose, then its infobox facts on one line."""
    t = re.sub(r"<!--.*?-->", "", t, flags=re.S)
    t = re.sub(r"<ref[^>]*/>|<ref.*?</ref>", "", t, flags=re.S)
    box = []
    t = _strip_templates(t, box)
    t = re.sub(r"\[\[(?:File|Image|Category|[a-z]{2}(?:-[a-z]+)?):[^\]]*\]\]", "", t, flags=re.I)
    t = re.sub(r"\[\[[^\]|]*\|([^\]]*)\]\]", r"\1", t)
    t = re.sub(r"\[\[([^\]]*)\]\]", r"\1", t)
    t = re.sub(r"\[https?://\S+ ([^\]]*)\]", r"\1", t)
    t = re.sub(r"'{2,}", "", t)
    t = re.sub(r"<[^>]+>", "", t)
    t = re.sub(r"^\{\|.*$|^\|\}.*$|^\|-.*$", "", t, flags=re.M)
    t = re.sub(r"^[|!]\s*", "  ", t, flags=re.M)
    t = re.sub(r"^(=+)\s*(.*?)\s*\1\s*$", r"\n## \2", t, flags=re.M)
    t = re.sub(r"__[A-Z]+__", "", t)
    t = re.sub(r"\n{3,}", "\n\n", html.unescape(t)).strip()
    if box:
        t += "\n[facts] " + "; ".join(dict.fromkeys(box))[:600]
    return t


def page_texts(base, titles, raw=False):
    """{title: text} for up to 50 titles per request, following redirects.
    `raw` keeps the wikitext as written: right for pages whose substance is in
    nested templates (an event's choices, a table of outcomes)."""
    out = {}
    for i in range(0, len(titles), 50):
        chunk = titles[i:i + 50]
        d = api(base, action="query", prop="revisions", rvprop="content", rvslots="main",
                titles="|".join(chunk), redirects="1")
        if not d:
            continue
        q = d.get("query", {})
        back = {r["to"]: r["from"] for r in q.get("redirects", [])}
        norm = {n["to"]: n["from"] for n in q.get("normalized", [])}
        for p in q.get("pages", {}).values():
            revs = p.get("revisions")
            if not revs:
                continue
            slot = revs[0].get("slots", {}).get("main", revs[0])
            text = slot.get("*") or slot.get("content") or ""
            t = p["title"]
            t = back.get(t, t)
            out[norm.get(t, t)] = _raw(text) if raw else wikitext_to_text(text)
    return out


def _raw(text):
    """Wikitext minus the parts that are never content: comments, file links
    and galleries."""
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    text = re.sub(r"<gallery.*?</gallery>", "", text, flags=re.S | re.I)
    text = re.sub(r"\[\[(?:File|Image):[^\]]*\]\]", "", text, flags=re.I)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def category_members(base, category, limit=500):
    cat = category if category.lower().startswith("category:") else "Category:" + category
    # Articles only (namespace 0): a wiki's Data:, Template: and Modding: pages
    # sit in the same categories and are machinery, not content.
    out, params = [], dict(action="query", list="categorymembers", cmtitle=cat, cmtype="page",
                           cmnamespace="0", cmlimit=str(min(limit, 500)))
    while len(out) < limit:
        d = api(base, **params)
        if not d:
            break
        out += [m["title"] for m in d.get("query", {}).get("categorymembers", [])]
        cont = d.get("continue", {}).get("cmcontinue")
        if not cont:
            break
        params["cmcontinue"] = cont
    return out[:limit]


# ── commands ────────────────────────────────────────────────────────────────

def order_games(kind, games, wikis, ledger, include_unowned=False, include_done=False):
    """Games to research next for `kind`: owned first, then the best-documented
    (the most wiki pages in that kind's categories), then the oldest."""
    out = []
    for g, info in games.items():
        if not include_done and researched(ledger, g, kind):
            continue
        if not include_unowned and not info["owned"]:
            continue
        out.append(g)
    out.sort(key=lambda g: (not games[g]["owned"], -wiki_score(wikis.get(g), kind),
                            str(games[g]["year"] or "9999"), g.lower()))
    return out


def cmd_next(args):
    games, wikis, ledger = catalog(), load_json(WIKIS, {}), load_ledger()
    todo = order_games(args.kind, games, wikis, ledger, include_unowned=args.unowned)
    if not todo and not args.unowned:
        todo = order_games(args.kind, games, wikis, ledger, include_unowned=True)
    for g in todo[:args.n]:
        w = wikis.get(g) or {}
        where = (f"{w.get('sitename')} — {w.get('counts', {}).get(args.kind, 0)} pages in {args.kind} "
                 f"categories" if w.get("wiki") else ("no wiki found" if g in wikis else "wiki not looked for"))
        print(f"{g} ({games[g]['year']}, {games[g]['type']}{', owned' if games[g]['owned'] else ''}): {where}")
    if not todo:
        print(f"every game has been researched for {args.kind}")


HAND_FIELDS = ("skip", "note")


def cmd_wikis(args):
    games, wikis = catalog(), load_json(WIKIS, {})
    if args.games:
        names = [g for g in args.games if g in games]
        for g in set(args.games) - set(names):
            print(f"not a game on the sheet: {g}")
    elif args.refresh:
        names = list(games)
    else:
        names = [g for g in games if g not in wikis]
    names = [g for g in names if (wikis.get(g) or {}).get("set") != "by hand"]
    if args.blurbs:
        # Records found before main pages were read: read them now, no rediscovery.
        todo = [g for g, w in wikis.items() if w.get("api") and w.get("set") != "by hand"]
        print(f"reading main pages for {len(todo)} wiki(s)", flush=True)
        with cf.ThreadPoolExecutor(args.threads) as ex:
            for g, raw in zip(todo, ex.map(lambda g: main_page(wikis[g]["api"]), todo)):
                wikis[g]["blurb"] = blurb_of(raw)
                wikis[g]["signal"] = signal_of(games[g]["year"], raw)
        names = []
    print(f"looking for wikis for {len(names)} game(s)", flush=True)
    done = 0
    with cf.ThreadPoolExecutor(args.threads) as ex:
        futs = {ex.submit(discover, g, games[g]["year"]): g for g in names}
        for f in cf.as_completed(futs):
            g = futs[f]
            try:
                old = wikis.get(g) or {}
                wikis[g] = f.result()
                # Fields only a person writes survive a rediscovery.
                wikis[g].update({k: old[k] for k in HAND_FIELDS if k in old})
            except Exception as e:  # one bad wiki never stops the pass
                print(f"  {g}: failed ({e})")
                continue
            done += 1
            if done % 25 == 0:
                save_json(WIKIS, wikis)
                print(f"  {done}/{len(names)}", flush=True)
    # The match rule can change after a wiki was found; it reads only the
    # stored sitename, so every record is re-judged by the current one.
    for g, w in wikis.items():
        if not w.get("sitename") or w.get("set") == "by hand":
            continue
        if g in KNOWN_WIKIS and w.get("api") == KNOWN_WIKIS[g]:
            w["match"] = "sure"
            continue
        q = match_quality(g, w["sitename"]) or "check"
        if q == "sure" and "signal" in w and not looks_like_the_game(None, "", w.get("counts"), w["signal"]):
            q = "check"
        w["match"] = q
    save_json(WIKIS, wikis)
    found = sum(1 for g in names if (wikis.get(g) or {}).get("wiki"))
    check = sum(1 for g in names if (wikis.get(g) or {}).get("match") == "check")
    print(f"found a wiki for {found} of {len(names)} ({check} named differently from the game: "
          f"`match: check` in {os.path.relpath(WIKIS, ROOT)})")


def game_content(game):
    """{sheet: [names]} of what the game already gives the game, across every
    sheet a candidate can land on."""
    out = {}
    for sheet, (ncol, gcol) in TARGET_COLS.items():
        names = [r[ncol] for r in sheet_table(sheet) if r.get(gcol, "").lower() == game.lower() and r.get(ncol)]
        if names:
            out[sheet] = names
    return out


# What a row of each content sheet is, in a few words, for the inventory.
DESCRIBE = {
    "enemies": ("Goal", "Tag"), "bosses": ("Goal", "Tag"), "items": ("Description",), "trinkets": ("Description",),
    "cards": ("Description",), "weapons": ("Goal", "Passive"), "bags": ("Description",), "wands": ("Description",),
    "potions": ("On Player", "On Tile"), "scrolls": ("Description",), "events": ("Requirement",),
    "characters": ("Level Up",), "statuses": ("Type", "Combat"), "curses": ("Condition",),
    "locations": ("Goal", "Goal Effect"), "objects": ("Tag",), "tiles": ("Description",),
    "units": ("Type", "Description"), "evolutions": ("Requirement 1", "Requirement 2", "Outcome"),
}
EXTRA_SHEETS = {"abilities": ("Name", "Description"), "pills": ("Name", "Description")}
THIN = 10  # a content sheet with fewer rows than this is a gap worth filling


def name_index():
    """{normalised name: [where]} over every live content row and every
    candidate, so a wiki title the game already has is flagged wherever it
    came from — the same thing from another game is still a duplicate."""
    idx = {}

    def add(name, where):
        k = _title_norm(name)
        if k:
            idx.setdefault(k, []).append(where)
    for sheet, (ncol, gcol) in TARGET_COLS.items():
        for r in sheet_table(sheet):
            if r.get(ncol):
                add(r[ncol], f"{sheet}: {r.get(gcol) or '?'}")
    for sheet, (ncol, _) in EXTRA_SHEETS.items():
        for r in sheet_table(sheet):
            if r.get(ncol):
                add(r[ncol], sheet)
    for k in WIKI_KINDS:
        for r in read_rows(k):
            add(r.get(KINDS[k]["name"], ""), f"candidate {k}: {r.get('Game')}")
    return idx


def inventory_lines():
    """Everything the game already has, by system, with the thin areas named."""
    import collections
    out = ["# What the game already has", "",
           f"Written by `python3 tools/research.py inventory` on {TODAY} from tools/Roguelikes.xlsx. Read it "
           "before writing candidates: a candidate should fill a gap or plug into something below, and must "
           "not duplicate a row here (any game's version counts).", ""]
    sizes, source = {}, collections.Counter()
    for sheet, (ncol, gcol) in TARGET_COLS.items():
        rows = sheet_table(sheet)
        sizes[sheet] = len(rows)
        for r in rows:
            if r.get(gcol):
                source[r[gcol]] += 1
    for sheet, _ in EXTRA_SHEETS.items():
        sizes[sheet] = len(sheet_table(sheet))
    thin = sorted((n, s) for s, n in sizes.items() if n < THIN)
    out += ["## Thin areas (fewer than %d rows)" % THIN, ""] + [f"- {s}: {n}" for n, s in thin] + [
        "- keys: no content grants or spends one (spec §4)",
        "- shops: no shopkeepers authored, and the shelf sells only relics and loot (spec §14.5)", ""]
    out += ["## Where the content comes from", "",
            f"{sum(source.values())} rows from {len(source)} games; the top five: "
            + ", ".join(f"{g} {n}" for g, n in source.most_common(5)), ""]
    tags = collections.Counter(t.strip() for s in ("enemies", "bosses") for r in sheet_table(s)
                               for t in r.get("Tag", "").split(",") if t.strip())
    out += ["## Enemy tags (what a location, an ability or a goal can name)", "",
            ", ".join(f"{t} {n}" for t, n in tags.most_common()), ""]
    ltags = collections.Counter(t.strip() for s, col in (("items", "tags"), ("trinkets", "Tag"), ("cards", "Tags"),
                                                        ("weapons", "Tag"), ("potions", "Tags"), ("wands", "Tags"),
                                                        ("scrolls", "Tags"))
                                for r in sheet_table(s) for t in r.get(col, "").split(",") if t.strip())
    out += ["## Loot tags (families a new piece can join: evolutions and neighbours read them)", "",
            ", ".join(f"{t} {n}" for t, n in ltags.most_common()), ""]
    verbs = collections.Counter()
    for s in list(DESCRIBE) + ["pills", "bags"]:
        for r in sheet_table(s):
            for k, v in r.items():
                if "Effect" in k and v:
                    verbs.update(re.findall(r"(?:^|[;:>]\s*|else\s+|->\s*)([a-z_]{3,})", v))
    out += ["## The effect language in use (verbs and triggers, by how often)", "",
            ", ".join(f"{k} {v}" for k, v in verbs.most_common()), ""]
    for sheet, cols in DESCRIBE.items():
        ncol, gcol = TARGET_COLS[sheet]
        rows = sheet_table(sheet)
        out += [f"## {sheet} ({len(rows)})", ""]
        out += [f"- {r[ncol]} ({r.get(gcol) or '?'}): " + " / ".join(r.get(c, "")[:90] for c in cols if r.get(c))
                for r in rows]
        out.append("")
    for sheet, (ncol, dcol) in EXTRA_SHEETS.items():
        rows = sheet_table(sheet)
        out += [f"## {sheet} ({len(rows)})", ""] + [f"- {r[ncol]}: {r.get(dcol, '')[:110]}" for r in rows] + [""]
    return out


def cmd_inventory(args):
    text = "\n".join(inventory_lines())
    os.makedirs(WORK, exist_ok=True)
    path = os.path.join(WORK, "inventory.md")
    with open(path, "w", encoding="utf8") as fh:
        fh.write(text)
    print(text if not args.quiet else f"wrote {os.path.relpath(path, ROOT)}")


def weapon_leads(base, members):
    """The brief's standing section: where this wiki keeps its weapons and its
    evolutions (synergies, fusions, recipes), found live by name so a wiki
    whose loot categories bury them still shows them."""
    cats = count_categories(base)
    allc = {c for v in cats.values() for c, _ in v}
    wcats = sorted((c for c in allc if WEAPON_CATS.search(c)), key=len)[:4]
    ecats = sorted((c for c in allc if EVOLUTION_CATS.search(c)  # "Evolved Bosses" is a monster list
                    and not re.search(CATEGORY["goals"], c, re.I)), key=len)[:4]
    have = {}
    d = api(base, action="query", titles="|".join(WEAPON_PAGES + EVOLUTION_PAGES), redirects="1")
    if d:
        q = d.get("query", {})
        back = {r["to"]: r["from"] for r in q.get("redirects", [])}
        have = {back.get(p["title"], p["title"]) for p in q.get("pages", {}).values()
                if "missing" not in p and "invalid" not in p}
    out = ["## Weapons and evolutions (asked of every game)", "",
           "A weapon is swung or fired: `Sheet` `weapons`, with `Extra` Aim, Area, Type, Goal, Charge. An "
           "evolution turns a weapon into a better one when a named piece or a tagged family is held: `Sheet` "
           "`evolutions`, Name = what it becomes, `Extra` Requirement 1 (the weapon), Requirement 2, Outcome "
           "(docs/research.md §6.2).", ""]
    for label, found, pages in (("Weapons", wcats, WEAPON_PAGES), ("Evolutions", ecats, EVOLUTION_PAGES)):
        lp = [p for p in pages if p in have]
        out += [f"### {label}", "", "List articles: " + (", ".join(lp) or "none"), ""]
        for c in found[:2]:
            m = category_members(base, c, members)
            if m:
                out += [f"Category:{c} ({len(m)}{'+' if len(m) >= members else ''}): " + ", ".join(m), ""]
        if not lp and not found:
            out += [f"- no {label.lower()} category or page: say so in the ledger note", ""]
    return out


def cmd_brief(args):
    games = catalog()
    if args.game not in games:
        sys.exit(f"not a game on the sheet: {args.game}")
    info, wikis, ledger = games[args.game], load_json(WIKIS, {}), load_ledger()
    kinds = [args.kind] if args.kind else WIKI_KINDS
    w = wikis.get(args.game) or {}
    lines = [f"# {args.game}", "",
             f"- {info['year']}, {info['type']}, {'owned' if info['owned'] else 'not owned'}",
             f"- Steam: {info['steam'] or 'none'}",
             f"- Wiki: {w.get('wiki') or ('none found' if args.game in wikis else 'not looked for: run `wikis`')}"
             + (f" ({w.get('sitename')}, {w.get('articles')} articles, match: {w.get('match')})"
                if w.get("wiki") else ""),
             "- Researched: " + (", ".join(f"{k} {v}" for k, v in sorted(ledger.get(args.game, {}).items()))
                                 or "nothing yet"), ""]
    have = game_content(args.game)
    lines += ["## Already in this game", ""] + ([f"- {s}: {', '.join(n)}" for s, n in have.items()]
                                                 or ["- nothing"]) + [""]
    cands = []
    for k in KINDS:
        for r in read_rows(k):
            if r.get("Game") == args.game or (k == "connections" and args.game in (r["Influencer"],
                                                                                  r["Influencee"])):
                if k == "goals":  # in full: a redo pass checks each of these against the wiki
                    cands.append(f"- goals: {r['Name']} ({r['Sheet']}, {r['Difficulty']}, {r['Goal Type']}): "
                                 f"{r['Goal']} [{r['Confidence']}] {r['ID']}")
                else:
                    cands.append(f"- {k}: {r.get(KINDS[k]['name']) or r.get('Text', '')[:60]} [{r['Status']}]")
    lines += ["## Already a candidate", ""] + (cands or ["- nothing"]) + [""]
    base = w.get("api")
    known = name_index()

    def flag(title):
        hit = known.get(_title_norm(title))
        return f"{title} [HAVE: {'; '.join(hit[:2])}]" if hit else title
    lines += ["Titles marked [HAVE: …] are already in the game or a candidate, from any game. The whole "
              "inventory: `python3 tools/research.py inventory`.", ""]
    for k in kinds:
        cats = (w.get("categories") or {}).get(k, [])
        # A wiki that also covers other games (NetHack's covers a dozen variants)
        # names them in `skip` in wikis.json; their categories bury the game's.
        if w.get("skip"):
            cats = [c for c in cats if not re.search(w["skip"], c, re.I)]
        pages = (w.get("pages") or {}).get(k, [])
        lines += [f"## {k}: {KINDS[k]['title']}", ""]
        if not base:
            lines += ["- no wiki: work from the Steam page and what you know of the game", ""]
            continue
        if pages:
            lines += ["List articles: " + ", ".join(pages), ""]
        if not cats:
            lines += ["- no category for this kind" + ("; read the list articles above" if pages else ""), ""]
            continue
        shown = 0
        for c in cats:
            if shown >= args.categories:
                break
            members = category_members(base, c, args.members)
            if not members:  # every page in it was machinery (Data:, Template:)
                continue
            shown += 1
            lines += [f"### Category:{c} ({len(members)}{'+' if len(members) >= args.members else ''})", "",
                      ", ".join(flag(m) for m in members), ""]
    if base and (not args.kind or args.kind == "loot"):
        lines += weapon_leads(base, args.members)
    lines += ["Read pages: `python3 tools/research.py page \"%s\" \"<title>\" ...` or `--category \"<name>\"`."
              % args.game, "Rules for each kind's rows: docs/research.md.", ""]
    text = "\n".join(lines)
    os.makedirs(os.path.join(WORK, "briefs"), exist_ok=True)
    path = os.path.join(WORK, "briefs", slug(args.game) + ".md")
    with open(path, "w", encoding="utf8") as fh:
        fh.write(text)
    print(text if not args.quiet else f"wrote {os.path.relpath(path, ROOT)}")


def cmd_page(args):
    wikis = load_json(WIKIS, {})
    w = wikis.get(args.game) or {}
    base = args.api or w.get("api")
    if not base:
        sys.exit(f"no wiki recorded for {args.game}: run `wikis \"{args.game}\"`, or pass --api")
    titles = list(args.titles)
    if args.category:
        titles += category_members(base, args.category, args.limit)
    texts = page_texts(base, titles, raw=args.raw)
    for t in titles:
        body = texts.get(t)
        if body is None:
            print(f"=== {t}: (no page)\n")
            continue
        cut = body if args.chars <= 0 else body[:args.chars] + ("…" if len(body) > args.chars else "")
        print(f"=== {t}\n{cut}\n")


def cmd_mark(args):
    games = catalog()
    bad = [g for g in args.games if g not in games]
    if bad:
        sys.exit("not games on the sheet: " + ", ".join(bad))
    mark(args.games, args.kind, args.note)
    print(f"marked {len(args.games)} game(s) researched for {args.kind}")


def refresh_status(kind, rows, idx=None, pairs=None, names=None, verbose=False):
    """Move rows the owner has put in the sheet to `on sheet`, and rows waiting
    for a game the sheet now has back to `to review`. Returns what moved."""
    moved = []
    for r in rows:
        before = r["Status"]
        if kind == "connections":
            a, b = r["Influencer"], r["Influencee"]
            if not (a and b):
                continue
            if before in ("to review", "waiting for game row") and (a.lower(), b.lower()) in pairs:
                r["Status"] = "on sheet"
            elif before == "waiting for game row" and a in names and b in names:
                r["Status"] = "to review"
        else:
            sheet = target_sheet(kind, r)
            name = r.get(KINDS[kind]["name"], "").lower()
            if before == "to review" and name and name in idx.get(sheet, {}):
                r["Status"] = "on sheet"
        if r["Status"] != before:
            moved.append((r, before))
            if verbose:
                print(f"  {kind}: {r.get(KINDS[kind]['name'])} ({r.get('Game') or r.get('Influencer')}) "
                      f"{before} -> {r['Status']}")
    return moved


# A candidate whose Status says the owner should NOT add it, found on the sheet
# anyway: the owner decided otherwise, or added it for another reason. Never
# changed by itself; listed for a person to settle.
SETTLED_NO = ("not an influence", "nothing found", "lead")


def sheet_conflicts(kind, rows, idx=None, pairs=None):
    """Rows the sheet disagrees with, as (row, why) — read-only. For connections:
    a pair the owner added the OTHER way round (which direction is right?), and
    a pair on the sheet whose row says it isn't an influence or wasn't found."""
    out = []
    for r in rows:
        if kind == "connections":
            a, b = r["Influencer"].lower(), r["Influencee"].lower()
            if not (a and b):
                continue
            if r["Status"] in SETTLED_NO and (a, b) in pairs:
                out.append((r, f"is on the connections sheet, but this row says `{r['Status']}`"))
            elif r["Status"] in ("to review", "waiting for game row") and (b, a) in pairs \
                    and (a, b) not in pairs:
                out.append((r, "is on the sheet the OTHER way round"))
        else:
            name = r.get(KINDS[kind]["name"], "").lower()
            sheet = target_sheet(kind, r)
            on = (idx or {}).get(sheet, {})
            if r["Status"] == "to review" and name in on and on[name].lower() != r.get("Game", "").lower():
                out.append((r, f"has a namesake on the {sheet} sheet from {on[name] or 'another game'}"))
    return out


def label(kind, r):
    if kind == "connections":
        return f"{r['Influencer']} → {r['Influencee']}" if r["Influencer"] else r["Text"][:60]
    return f"{r.get(KINDS[kind]['name'])} ({r.get('Game')})"


def sheet_report():
    """What `build` would change, and what it would flag, without writing:
    (would move, conflicts). `check` prints it, so CI says when the owner has
    added rows the research sheet doesn't know about yet."""
    idx, pairs, names = on_sheet_index(), connection_pairs(), set(catalog())
    moves, conflicts = [], []
    for kind in KINDS:
        rows = [dict(r) for r in read_rows(kind)]
        moves += [(kind, r, before) for r, before in refresh_status(kind, rows, idx, pairs, names)]
        conflicts += [(kind, r, why) for r, why in sheet_conflicts(kind, rows, idx, pairs)]
    return moves, conflicts


def assign_ids(kind, rows):
    seen = {r["ID"] for r in rows if r.get("ID")}
    prefix = kind[:3]
    for r in rows:
        if r.get("ID"):
            continue
        if kind == "connections":
            base = (f"{prefix}-{slug(r['Influencer'])}--{slug(r['Influencee'])}" if r["Influencer"]
                    else f"{prefix}-{slug(r['Status'])}-{slug(r.get('Game') or r['Text'][:40])}")
        else:
            base = f"{prefix}-{slug(r.get('Game', ''))}-{slug(r.get(KINDS[kind]['name'], ''))}"
        rid, n = base, 2
        while rid in seen:
            rid, n = f"{base}-{n}", n + 1
        r["ID"] = rid
        seen.add(rid)


def extra_of(row):
    """A loot row's `Extra` as a dict: `Aim=front; Goal=Be a hero` -> {Aim: front, Goal: Be a hero}."""
    out = {}
    for part in row.get("Extra", "").split(";"):
        if "=" in part:
            k, v = part.split("=", 1)
            out[k.strip()] = v.strip()
    return out


def check_loot_extra(rows, where_of):
    """Weapons and evolutions carry their sheet's own columns in `Extra`; a
    missing one would paste as a blank the generator refuses."""
    bad = []
    weapons = {r["Name"].lower() for r in sheet_table("weapons") if r.get("Name")}
    weapons |= {r["Name"].lower() for r in rows if r.get("Sheet") == "weapons"}
    for r in rows:
        sheet = r.get("Sheet")
        if sheet not in EXTRA_KEYS:
            continue
        ex, where = extra_of(r), where_of(r)
        missing = [k for k in EXTRA_KEYS[sheet] if not ex.get(k)]
        if missing:
            bad.append(f"{where}: {r['Name']} ({sheet}) has no {', '.join(missing)} in Extra")
        if sheet == "weapons" and ex.get("Aim") and ex["Aim"] not in WEAPON_AIMS \
                and not re.fullmatch(r"column \d", ex["Aim"]):
            bad.append(f"{where}: {r['Name']} Aim {ex['Aim']!r} is not one of {sorted(WEAPON_AIMS)} or `column N`")
        if sheet == "evolutions":
            if ex.get("Outcome") and ex["Outcome"] not in EVOLUTION_OUTCOMES:
                bad.append(f"{where}: {r['Name']} Outcome {ex['Outcome']!r} should be Consume All or Consume None")
            if ex.get("Requirement 1") and ex["Requirement 1"].lower() not in weapons:
                bad.append(f"{where}: {r['Name']} evolves from {ex['Requirement 1']!r}, which is neither a live "
                           "weapon nor a weapon candidate (Requirement 1 is always the weapon that turns)")
    return bad


def check(verbose=True):
    """Problems in the store, as a list of strings. Read-only."""
    bad = []
    games = catalog()
    for kind, k in KINDS.items():
        path = csv_path(kind)
        if not os.path.exists(path):
            continue
        with open(path, newline="", encoding="utf-8") as fh:
            head = next(csv.reader(fh), [])
        if head != k["columns"]:
            bad.append(f"{kind}.csv: header is {head}, expected {k['columns']}")
            continue
        ids, keys = set(), set()
        for n, r in enumerate(read_rows(kind), 2):
            where = f"{kind}.csv line {n}"
            if not r["ID"]:
                pass  # a new row: `build` assigns its ID, and `check` notes it via book_drift
            elif r["ID"] in ids:
                bad.append(f"{where}: ID {r['ID']} is used twice")
            ids.add(r["ID"])
            if r["Status"] not in STATUSES:
                bad.append(f"{where}: Status {r['Status']!r} is not one of {sorted(STATUSES)}")
            if r["Owner"].lower() not in OWNER:
                bad.append(f"{where}: Owner {r['Owner']!r} should be yes, no or empty")
            if kind == "connections":
                continue
            if r.get("Confidence", "") not in CONFIDENCE:
                bad.append(f"{where}: Confidence {r['Confidence']!r} should be ok, check or empty")
            if r["Game"] not in games:
                bad.append(f"{where}: Game {r['Game']!r} is not spelled as the games sheet spells it")
            if "sheets" in k and r.get("Sheet") not in k["sheets"]:
                bad.append(f"{where}: Sheet {r.get('Sheet')!r} should be one of {k['sheets']}")
            for clone in filter(None, (c.strip() for c in r.get("Clones", "").split(";"))):
                head = clone.split(": ", 1)[0]
                game = head[head.rfind(" (") + 2:-1] if head.endswith(")") and " (" in head else ""
                if game not in games:
                    bad.append(f"{where}: clone {clone[:60]!r} should read `Name (Game): how it differs`, "
                               f"with the game spelled as the games sheet spells it")
            key = (target_sheet(kind, r), r.get(k["name"], "").lower())
            if not key[1]:
                bad.append(f"{where}: no {k['name']}")
            elif key in keys:
                bad.append(f"{where}: {r[k['name']]} is a candidate twice for {key[0]}")
            keys.add(key)
        if kind == "loot":
            rows = read_rows(kind)
            line = {id(r): n for n, r in enumerate(rows, 2)}
            bad += check_loot_extra(rows, lambda r: f"loot.csv line {line[id(r)]}")
    ledger = load_ledger()
    for g, kinds in ledger.items():
        if g not in games:
            bad.append(f"ledger.json: {g!r} is not a game on the sheet (renamed? move its entry)")
        for kd in kinds:
            if kd not in LEDGER_KINDS:
                bad.append(f"ledger.json: {g!r} has unknown kind {kd!r}")
    for g, w in load_json(WIKIS, {}).items():
        if g not in games:
            bad.append(f"wikis.json: {g!r} is not a game on the sheet")
    return bad


def cmd_check(args):
    bad = check()
    stale = book_drift()
    moves, conflicts = sheet_report()
    if moves:
        stale.append(f"{len(moves)} candidate(s) are on Roguelikes.xlsx now and still say otherwise "
                     f"(e.g. {label(moves[0][0], moves[0][1])}): `sync` updates them")
    for line in bad:
        print(line)
    for line in stale:
        print("note: " + line)
    for k, r, why in conflicts:
        print(f"look: {k} {label(k, r)} {why}")
    if bad:
        print(f"{len(bad)} problem(s)")
        return 1
    print("ok — every research CSV has its columns, a unique ID, a known Status and Owner, and names real games"
          + ("; Research.xlsx and the CSVs differ (see notes): `sync` if the owner has uploaded the workbook, "
             "otherwise `build`" if stale else ""))
    return 0


# ── the workbook ────────────────────────────────────────────────────────────

WIDTHS = {"Name": 24, "Event": 24, "Game": 24, "Influencer": 24, "Influencee": 24, "Goal": 36,
          "Description": 44, "Why it fits": 44, "Why this pairing": 44, "Text": 90, "Source": 30,
          "Prompt": 44, "Owner": 8, "Owner Notes": 30, "ID": 14, "Status": 13, "Heading": 24,
          "Level Up": 34, "What it is there": 40, "Builds on": 22, "Needs": 22, "Clones": 30, "On Player": 30, "On Enemy": 30, "Choices": 40,
          "Extra": 22, "Tag": 14, "Tags": 14, "Sheet": 10, "Confidence": 10, "Effect": 24}
WRAP = {"Description", "Why it fits", "Why this pairing", "Text", "Prompt", "What it is there", "Choices", "Clones"}


def book_rows(wb, kind):
    head, rows = rb.read_sheet(wb, kind)
    return head, rows


def book_drift():
    """Rows that differ between Research.xlsx and the CSVs, read-only. Either
    the owner edited the workbook (`sync` brings that in) or a CSV changed
    without a `build`."""
    if not os.path.exists(rb.BOOK):
        return ["Research.xlsx does not exist yet: run `build`"]
    wb = rb.open_book()
    notes = []
    for kind in KINDS:
        rows = read_rows(kind)
        if not rows and kind not in wb.sheetnames:
            continue
        _, brows = book_rows(wb, kind)
        cols = KINDS[kind]["columns"]
        by_id = {r.get("ID"): r for r in brows}
        fresh = sum(1 for r in rows if not r["ID"])
        if fresh:
            notes.append(f"{kind}: {fresh} new row(s) with no ID yet; `build` assigns them")
        diff = sum(1 for r in rows if r["ID"] and (r["ID"] not in by_id
                   or [by_id[r["ID"]].get(c, "") for c in cols] != [r[c] for c in cols]))
        extra = sum(1 for r in brows if not r.get("ID"))
        if diff or extra:
            notes.append(f"{kind}: {diff} row(s) differ from the CSV" + (f", {extra} new in the workbook"
                                                                          if extra else ""))
    return notes


def write_kind_sheet(wb, kind, rows):
    cols = KINDS[kind]["columns"]
    ws = rb.replace_sheet(wb, kind)
    ws.append(cols + ["_base"])
    for r in rows:
        vals = [r.get(c, "") for c in cols]
        ws.append([rb.cell_value(v) for v in vals] + [rb.row_hash(vals)])
    rb.style_table(ws, WIDTHS, owner_col="Owner", wrap=WRAP)
    letter = openpyxl.utils.get_column_letter(len(cols) + 1)
    ws.column_dimensions[letter].hidden = True


def write_status_sheet(wb, games, ledger, wikis, conflicts=()):
    ws = rb.replace_sheet(wb, "status")
    owned = [g for g in games if games[g]["owned"]]
    ws.append(["Kind", "Games researched", "…of the owned", "Candidates", "To review", "Owner: yes",
               "Owner: no", "On the sheet now", "Next up"])
    for kind in KINDS:
        rows = read_rows(kind)
        done = sum(1 for g in games if researched(ledger, g, kind))
        done_owned = sum(1 for g in owned if researched(ledger, g, kind))
        nxt = order_games(kind, games, wikis, ledger)[:5] if kind in WIKI_KINDS else []
        ws.append([kind, f"{done} / {len(games)}", f"{done_owned} / {len(owned)}", len(rows),
                   sum(r["Status"] == "to review" for r in rows), sum(r["Owner"] == "yes" for r in rows),
                   sum(r["Owner"] == "no" for r in rows), sum(r["Status"] == "on sheet" for r in rows),
                   ", ".join(nxt)])
    _, tags = rb.read_sheet(wb, "tag suggestions")
    if tags:
        ws.append(["tags", "every game with a Steam page", "", len(tags), sum(not t.get("Owner") for t in tags),
                   sum(t.get("Owner") == "yes" for t in tags), sum(t.get("Owner") == "no" for t in tags), "",
                   "rerun tools/tag_research.py"])
    if conflicts:
        ws.append([])
        ws.append(["Needs a look: the sheet and these candidates disagree. Settle each one (fix the "
                   "candidate's Status, or the sheet) and the line goes away."])
        for kind, r, why in conflicts:
            ws.append([kind, label(kind, r), why, r["ID"]])
    have = sum(1 for w in wikis.values() if w.get("wiki"))
    ws.append([])
    ws.append([f"Wikis: {have} found of {len(wikis)} games looked for ({len(games)} on the sheet). "
               f"Written {TODAY} by `python3 tools/research.py build`."])
    rb.style_table(ws, {"Kind": 13, "Games researched": 18, "…of the owned": 14, "Next up": 90})


ABOUT = [
    "Research — every candidate the project has found, for the owner to review. Nothing here is in Roguelikes.xlsx.",
    "",
    "HOW TO REVIEW",
    "  Each sheet is one kind of research. Write yes or no in the yellow Owner column; add anything to Owner Notes.",
    "  You can also correct any other cell (a goal's wording, a rarity): the next `sync` keeps your edit.",
    "  Upload the workbook back to tools/Research.xlsx. A Claude session then runs `python3 tools/research.py sync`.",
    "  A `yes` row waits there until it is in Roguelikes.xlsx; its Status then turns to `on sheet` by itself.",
    "  Don't delete rows: write no instead, so the same idea isn't suggested again.",
    "  Clones lists ideas that overlapped a row and lost to it: Name (Game): how it differs. Prefer one? Say so in Owner Notes.",
    "",
    "THE SHEETS",
    "  status       what has been researched, for which games, and what is next",
    "  connections  influences between chart games: the developer's own words, with the link (Text, Source)",
    "  goals        enemies and bosses with a goal, in the enemies/bosses sheets' own columns",
    "  loot         items, trinkets, cards, weapons, bags, wands, potions and scrolls (Sheet says which)",
    "  events       events in the events sheet's own columns; the Effect cells are left for you",
    "  characters   playable characters in the characters sheet's columns",
    "  statuses     statuses and curses (Sheet says which)",
    "  locations    locations and objects (Sheet says which)",
    "  tag suggestions / new tag ideas / vocabulary   from tools/tag_research.py, as before",
    "",
    "COLUMNS EVERY CANDIDATE SHEET HAS",
    "  Why it fits   the case for the row, and how it would play here",
    "  Source        the wiki page, store page or interview it came from",
    "  Confidence    ok, or check: written from a summary, or a detail nobody could confirm",
    "  Status        " + "; ".join(f"{k} = {v}" for k, v in STATUSES.items()),
    "  ID            how the row is matched when you sort or filter. Leave it alone.",
    "",
    "Effect cells are deliberately blank: a mechanic is the owner's to write.",
    "How the research is done: docs/research.md.",
]


def write_about_sheet(wb):
    ws = rb.replace_sheet(wb, "about")
    for line in ABOUT:
        ws.append([line])
    ws.column_dimensions["A"].width = 140
    ws["A1"].font = rb.BOLD


def build(verbose=False):
    games, ledger, wikis = catalog(), load_ledger(), load_json(WIKIS, {})
    idx, pairs, names = on_sheet_index(), connection_pairs(), set(games)
    wb = rb.open_book()
    built, conflicts = [], []
    for kind in KINDS:
        rows = read_rows(kind)
        if not rows and not os.path.exists(csv_path(kind)):
            continue
        assign_ids(kind, rows)
        refresh_status(kind, rows, idx, pairs, names, verbose)
        conflicts += [(kind, r, why) for r, why in sheet_conflicts(kind, rows, idx, pairs)]
        write_rows(kind, rows)
        write_kind_sheet(wb, kind, rows)
        built += [(kind, r["ID"]) for r in rows]
    write_about_sheet(wb)
    write_status_sheet(wb, games, ledger, wikis, conflicts)
    for kind, r, why in conflicts:
        print(f"  look: {kind} {label(kind, r)} {why}")
    ws = rb.replace_sheet(wb, "_built")
    ws.append(["kind", "id"])
    for row in built:
        ws.append(list(row))
    rb.save(wb)
    return len(built)


def cmd_build(args):
    n = build(verbose=True)
    print(f"wrote {n} candidate rows to {os.path.relpath(rb.BOOK, ROOT)}")


def built_as(row, cols, base):
    """Whether `row` is the row whose hash is `base`. A workbook built before a
    LATE_COLUMNS column existed hashed its rows without it, so a row that leaves
    those columns empty also matches the hash taken without them."""
    if rb.row_hash([row.get(c, "") for c in cols]) == base:
        return True
    late = [c for c in cols if c in LATE_COLUMNS]
    return bool(late) and not any(row.get(c, "") for c in late) and \
        rb.row_hash([row.get(c, "") for c in cols if c not in LATE_COLUMNS]) == base


def cmd_sync(args):
    """Bring the owner's edits in Research.xlsx into the CSVs, then rebuild.

    Every row in the workbook carries a hidden `_base`, the hash of the row as
    it was built. A row whose cells no longer hash to it was edited by the
    owner; a CSV row that no longer hashes to it was edited by a session since.
    One side changed: that side wins. Both changed: the owner's Owner and Owner
    Notes are taken, the session's other cells are kept, and the clash is
    printed so a person can settle it. A row the owner added (no ID) is taken
    as a new candidate; a row they deleted is reported, never dropped.
    """
    if not os.path.exists(rb.BOOK):
        sys.exit("no Research.xlsx to sync from: run `build`")
    wb = rb.open_book()
    _, built = rb.read_sheet(wb, "_built")
    built_ids = {(b["kind"], b["id"]) for b in built}
    taken = clashes = added = 0
    for kind in KINDS:
        if kind not in wb.sheetnames:
            continue
        cols = KINDS[kind]["columns"]
        rows = read_rows(kind)
        by_id = {r["ID"]: r for r in rows}
        _, brows = rb.read_sheet(wb, kind)
        seen = set()
        for b in brows:
            vals = [b.get(c, "") for c in cols]
            rid = b.get("ID", "")
            if not rid:
                if any(vals):
                    rows.append({c: b.get(c, "") for c in cols} | {"Status": b.get("Status") or "to review"})
                    added += 1
                continue
            seen.add(rid)
            r = by_id.get(rid)
            if r is None:
                print(f"  {kind}: {rid} is in the workbook but not the CSV; left out (was it removed on purpose?)")
                continue
            base = b.get("_base", "")
            owner_changed = not built_as(b, cols, base)
            csv_changed = not built_as(r, cols, base)
            if not owner_changed:
                continue
            if not csv_changed:
                for c in cols:
                    if r[c] != b.get(c, ""):
                        r[c] = b.get(c, "")
                taken += 1
            else:
                for c in ("Owner", "Owner Notes"):
                    r[c] = b.get(c, "")
                diff = [c for c in cols if c not in ("Owner", "Owner Notes") and r[c] != b.get(c, "")]
                if diff:
                    clashes += 1
                    print(f"  CLASH {kind} {rid}: both sides changed {', '.join(diff)}; kept the CSV's. "
                          f"Workbook had: " + "; ".join(f"{c}={b.get(c, '')!r}" for c in diff))
                else:
                    taken += 1
        for r in rows:
            if (kind, r["ID"]) in built_ids and r["ID"] not in seen:
                print(f"  {kind}: {r['ID']} ({r.get(KINDS[kind]['name'])}) was deleted from the workbook; "
                      f"kept. Write `no` in Owner to turn it down.")
        bad_owner = [r for r in rows if r["Owner"].lower() not in OWNER]
        for r in rows:
            r["Owner"] = r["Owner"].lower() if r["Owner"].lower() in OWNER else r["Owner"]
        for r in bad_owner:
            print(f"  {kind}: {r['ID']} has Owner {r['Owner']!r}; `check` will flag it")
        write_rows(kind, rows)
    n = build(verbose=True)
    print(f"took {taken} edited row(s) and {added} new row(s) from the workbook, {clashes} clash(es); "
          f"rebuilt {n} rows")


def cmd_status(args):
    games, ledger, wikis = catalog(), load_ledger(), load_json(WIKIS, {})
    owned = [g for g in games if games[g]["owned"]]
    print("%-12s %12s %12s %10s %9s %5s %5s %8s" % ("kind", "researched", "of owned", "candidates", "to review",
                                                    "yes", "no", "on sheet"))
    for kind in KINDS:
        rows = read_rows(kind)
        print("%-12s %12s %12s %10d %9d %5d %5d %8d" % (
            kind, f"{sum(researched(ledger, g, kind) for g in games)}/{len(games)}",
            f"{sum(researched(ledger, g, kind) for g in owned)}/{len(owned)}", len(rows),
            sum(r["Status"] == "to review" for r in rows), sum(r["Owner"] == "yes" for r in rows),
            sum(r["Owner"] == "no" for r in rows), sum(r["Status"] == "on sheet" for r in rows)))
    looked = len(wikis)
    print(f"\nwikis: {sum(1 for w in wikis.values() if w.get('wiki'))} found, {looked} of {len(games)} games "
          f"looked for" + ("" if looked == len(games) else " (run `wikis` for the rest)"))
    if args.verbose:
        for kind in WIKI_KINDS:
            print(f"next for {kind}: " + ", ".join(order_games(kind, games, wikis, ledger)[:8]))
    drift = book_drift()
    for d in drift:
        print("Research.xlsx: " + d)


def cmd_new(args):
    """Everything a newly ported game gets: run the scripted passes over just
    the games with no ledger entry, write a brief for each, and list the hand
    work left. The ledger is NOT written here: each kind is marked with `mark`
    once its rows are written, so an interrupted session leaves the gaps visible."""
    games, ledger = catalog(), load_ledger()
    new = [g for g in games if g not in ledger]
    if args.games:
        new = [g for g in args.games if g in games]
    if not new:
        print("no new games: every game on the sheet has a ledger entry")
        return
    print(f"{len(new)} new game(s): {', '.join(new)}", flush=True)
    if not args.skip_influence:
        print("\n== connections: tools/influence_research.py new", flush=True)
        subprocess.call([sys.executable, os.path.join(TOOLS, "influence_research.py"), "new"])
    if not args.skip_tags:
        print("\n== tags: tools/tag_research.py", flush=True)
        subprocess.call([sys.executable, os.path.join(TOOLS, "tag_research.py")])
    print("\n== wikis", flush=True)
    cmd_wikis(argparse.Namespace(games=new, refresh=False, threads=4))
    for g in new:
        cmd_brief(argparse.Namespace(game=g, kind=None, members=60, categories=4, quiet=True))
    print("\nLeft to do, by hand, for each new game (docs/research.md says how):")
    print("  connections  read .influence_work/new_games.md, search the web; rows go in research/connections.csv")
    print("  tags         nothing: the suggestions are in Research.xlsx for the owner")
    print("  goals, loot, events, characters, statuses, locations")
    print("               read .research_work/briefs/<game>.md and the wiki pages it lists; rows go in research/<kind>.csv")
    print("  then:        python3 tools/research.py mark <kind> <games...> --note \"...\"  (once per kind)")
    print("               python3 tools/influence_research.py new --mark")
    print("               python3 tools/research.py check && python3 tools/research.py build")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("status", help="coverage per kind")
    p.add_argument("-v", "--verbose", action="store_true", help="also list the next games per kind")
    p.set_defaults(fn=cmd_status)

    p = sub.add_parser("next", help="the next games to research for a kind")
    p.add_argument("kind", choices=WIKI_KINDS)
    p.add_argument("-n", type=int, default=10)
    p.add_argument("--unowned", action="store_true", help="include games the owner does not own")
    p.set_defaults(fn=cmd_next)

    p = sub.add_parser("wikis", help="find each game's wiki and count its content categories")
    p.add_argument("games", nargs="*", help="just these games (default: every game not looked for yet)")
    p.add_argument("--refresh", action="store_true", help="look again for every game")
    p.add_argument("--blurbs", action="store_true", help="read every found wiki's main page again, no rediscovery")
    p.add_argument("--threads", type=int, default=6)
    p.set_defaults(fn=cmd_wikis)

    p = sub.add_parser("brief", help="everything about one game, for writing its candidates")
    p.add_argument("game")
    p.add_argument("--kind", choices=WIKI_KINDS)
    p.add_argument("--members", type=int, default=150, help="titles listed per category")
    p.add_argument("--categories", type=int, default=6, help="categories listed per kind")
    p.add_argument("--quiet", action="store_true", help="write the file, don't print it")
    p.set_defaults(fn=cmd_brief)

    p = sub.add_parser("page", help="wiki pages as plain text")
    p.add_argument("game")
    p.add_argument("titles", nargs="*")
    p.add_argument("--category", help="every page in this category")
    p.add_argument("--limit", type=int, default=200, help="pages read from --category")
    p.add_argument("--chars", type=int, default=2500, help="characters per page (0 = all)")
    p.add_argument("--api", help="a wiki's api.php, for a game with none recorded")
    p.add_argument("--raw", action="store_true", help="the wikitext as written (for pages built from templates)")
    p.set_defaults(fn=cmd_page)

    p = sub.add_parser("mark", help="record games as researched for a kind")
    p.add_argument("kind", choices=LEDGER_KINDS)
    p.add_argument("games", nargs="+")
    p.add_argument("--note", default="", help='e.g. "12 rows" or "nothing usable"')
    p.set_defaults(fn=cmd_mark)

    p = sub.add_parser("inventory", help="everything the game already has, by system, with its gaps")
    p.add_argument("--quiet", action="store_true", help="write .research_work/inventory.md, don't print it")
    p.set_defaults(fn=cmd_inventory)

    sub.add_parser("check", help="validate the CSVs, ledger and wikis.json").set_defaults(fn=cmd_check)
    sub.add_parser("build", help="write Research.xlsx from the CSVs").set_defaults(fn=cmd_build)
    sub.add_parser("sync", help="take the owner's edits from Research.xlsx, then build").set_defaults(fn=cmd_sync)

    p = sub.add_parser("new", help="every kind of research for the games just added")
    p.add_argument("games", nargs="*", help="these games instead of the ones with no ledger entry")
    p.add_argument("--skip-influence", action="store_true")
    p.add_argument("--skip-tags", action="store_true")
    p.set_defaults(fn=cmd_new)

    args = ap.parse_args()
    return args.fn(args) or 0


if __name__ == "__main__":
    sys.exit(main())
