#!/usr/bin/env python3
"""One-shot: Zoe's goal names the perfect, and Burn's way out says `game`.

TWO EDITS, in two different sheets, for two different reasons.

1. ZOE — `goals` sheet, the source (tools/apply_goals_sheet.py).

   "Beat a game without losing" -> "Perfect a game by beating it without losing"

   This is the goal the PERFECT-GAME FLAG hangs off. `Overworld2._means_perfected`
   reads the level-up condition's wording, because the condition has no
   machine-readable side, and the one thing that flag has already done is stop
   being set: the goals rewrite took Zoe's from "Perfect a Game" to "Beat a game
   without losing", and a bare `contains("perfect")` match went quietly dead.
   Putting the word back in the goal makes the prose say what the code is looking
   for, rather than leaving the code to recognise a paraphrase.

   It keeps "without losing" as well, so BOTH wordings in
   `Overworld2.PERFECTED_WORDINGS` match it and neither is load-bearing alone.

2. BURN — `statuses` sheet, `On Enemy`, which the `goals` sheet does not own.

   Gains "or instead beat a run  while skipping or trashing 4-X items/upgrades"
   Gains "or instead beat a game while skipping or trashing 4-X items/upgrades"

   `goals` carries only the PLAYER side of a status (Relation `modifies`), so
   pushing the sheet fixed Burn's `On Player` and left its `On Enemy` alone — and
   the two are a matched pair, the goal and the way out of it. One reading
   "beat a game" beside one reading "beat a run" is the same clause contradicting
   itself across two cells.

   This edits only Burn. The other five statuses use "run" in their `On Enemy`
   clauses too ("in the same run", "on a run where you beat X bosses"), and those
   are a different question — a paraphrase there is not a mismatch with their own
   player side — so they are left for a deliberate pass rather than swept up here.

Run once. `apply_goals_sheet.py --check` keeps edit 1 true; edit 2 is outside
that script's remit by design and is pinned by test_statuses.gd instead.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _xlsx_surgery import Workbook  # noqa: E402

import openpyxl  # noqa: E402

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
XLSX = os.path.join(ROOT, "tools", "Roguelikes.xlsx")

# (sheet, cell, expected old value, new value)
EDITS = [
    ("goals", "A99",
     "Beat a game without losing",
     "Perfect a game by beating it without losing"),
    ("statuses", "F2",
     'Gains "or instead beat a run while skipping or trashing 4-X items/upgrades"',
     'Gains "or instead beat a game while skipping or trashing 4-X items/upgrades"'),
]


def read(sheet: str, cell: str) -> str:
    """The value at `cell`, so the edit can check what it is replacing."""
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)[sheet]
    col = "".join(c for c in cell if c.isalpha())
    row = int("".join(c for c in cell if c.isdigit()))
    idx = 0
    for ch in col:
        idx = idx * 26 + (ord(ch) - ord("A") + 1)
    value = ws.cell(row=row, column=idx).value
    return "" if value is None else str(value).strip()


def main() -> int:
    planned = {}
    for sheet, cell, was, now in EDITS:
        found = read(sheet, cell)
        if found == now:
            print("%s!%s is already right — skipping" % (sheet, cell))
            continue
        # REFUSED rather than written blind. These are hand-picked cell
        # references; if the sheet has moved under them, writing anyway would put
        # a goal into whatever row is there now.
        if found != was:
            raise SystemExit(
                "%s!%s holds\n  %r\nbut this edit expects\n  %r\n"
                "The sheet moved — re-point the reference rather than "
                "overwriting it." % (sheet, cell, found, was))
        planned.setdefault(sheet, {})[cell] = now
        print("%s!%s\n  - %s\n  + %s" % (sheet, cell, was, now))

    if not planned:
        print("Nothing to do.")
        return 0

    with Workbook(XLSX) as wb:
        for sheet, edits in planned.items():
            wb.replace_cells(sheet, edits)

    print("\nNow run:  python3 tools/apply_goals_sheet.py"
          "  &&  python3 tools/generate_character2_tres.py"
          "  &&  python3 tools/generate_status_tres.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
