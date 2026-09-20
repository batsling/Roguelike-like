#!/usr/bin/env python3
"""One-shot: Speed's enemy clause gets a SUBJECT, so it composes into a goal.

An enemy-side `clause` is ANDed onto whatever goal the body is carrying
(`GameLoop2.goal_text_for`), so it has to be a phrase that survives being bolted
to the end of someone else's sentence. Two of the three do, because they name
what the condition is about:

    Dexterity   "you must beat {X} or all bosses without getting hit"
    Strength    "the difficulty must be increased {X} [time|times] or …"

Speed did not:

    Speed       "must be beaten in {…:hours} or less"

…which has no subject at all, and the composed line came out as

    Become undetectable and must be beaten in 2 hours or less
    Defeat 3 bugs and must be beaten in 2 hours or less

The thing that must be beaten is THE GAME, and the sentence never said so — it
reads as though the disguise, or the bugs, are what has to be beaten in two
hours. It is worst on an `any time` body (99 of the 134 goals), where the
sentence contains no game for a reader to recover the subject from; on a
`game beaten` body ("Beat a game without using magic and must be beaten in…")
you can just about reconstruct it, which is why it survived this long.

So: `the game must be beaten in …`, matching the other two clauses' shape.

The PLAYER side is left alone deliberately. `goal "beaten in {…} or less"`
renders as a standing row of its own — "If beaten in 2 hours or less, gain …" —
directly under the header `When beating a game:`, which supplies the subject the
phrase is missing. It is terse rather than broken, and changing it would put
"game" twice in three lines of the narrowest column on the page.

Run once, then tools/generate_status_tres.py.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _xlsx_surgery import Workbook  # noqa: E402

import openpyxl  # noqa: E402

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
XLSX = os.path.join(ROOT, "tools", "Roguelikes.xlsx")

# statuses!G5 — Speed's `On Enemy Effect`.
CELL = "G5"
WAS = 'clause "must be beaten in {1+(1/2)^(X-2):hours} or less"'
NOW = 'clause "the game must be beaten in {1+(1/2)^(X-2):hours} or less"'


def main() -> int:
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)["statuses"]
    found = ws.cell(row=5, column=7).value
    found = "" if found is None else str(found).strip()
    if found == NOW:
        print("statuses!%s is already right" % CELL)
        return 0
    if found != WAS:
        raise SystemExit(
            "statuses!%s holds\n  %r\nbut this edit expects\n  %r\n"
            "Re-point the reference rather than overwriting it." % (CELL, found, WAS))
    print("statuses!%s\n  - %s\n  + %s" % (CELL, WAS, NOW))
    with Workbook(XLSX) as wb:
        wb.replace_cells("statuses", {CELL: NOW})
    print("\nNow run:  python3 tools/generate_status_tres.py")
    return 0


if __name__ == "__main__":
    sys.exit(main())
