#!/usr/bin/env python3
"""`tools/Research.xlsx`: the one workbook every kind of research is reviewed in.

Two scripts write it — `research.py` (the candidate sheets, from the CSVs in
`research/`) and `tag_research.py` (the tag sheets, from Steam) — and each must
leave the other's sheets alone. Before this module, `tag_research.py` built the
workbook from nothing on every run, which was fine while it was the only writer.
Every writer now goes through `replace_sheet`, which swaps out the sheets it
owns and keeps everything else, in the order `ORDER` gives.

The workbook has no charts, so openpyxl is safe here. It is NOT safe on
`Roguelikes.xlsx` (see `_xlsx_surgery.py`); nothing in this module opens that.
"""

import hashlib
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BOOK = os.path.join(ROOT, "tools", "Research.xlsx")

# Sheet order in the workbook. A sheet not named here goes after these, so an
# owner's own scratch sheet survives a rebuild (at the end).
ORDER = ["about", "status", "connections", "goals", "loot", "events", "characters", "statuses",
         "locations", "tag suggestions", "new tag ideas", "vocabulary", "tags about", "_built"]

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
OWNER_FILL = PatternFill("solid", fgColor="FFF2CC")


def norm(v):
    """A cell as the CSV spells it: text, `1` not `1.0`, empty for nothing."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def cell_value(s):
    """The CSV's text as a cell: a whole number goes in as a number so it pastes
    into `Roguelikes.xlsx` as one; everything else stays text."""
    if s and s.isdigit() and len(s) < 10 and not (len(s) > 1 and s[0] == "0"):
        return int(s)
    return s if s else None


def row_hash(values):
    return hashlib.sha1("\x1f".join(values).encode("utf8")).hexdigest()[:12]


def open_book(path=BOOK):
    if os.path.exists(path):
        return openpyxl.load_workbook(path)
    wb = openpyxl.Workbook()
    wb.active.title = "about"
    return wb


def replace_sheet(wb, name):
    """An empty sheet called `name`, where `ORDER` says it goes, replacing any
    sheet of that name. Every other sheet is untouched."""
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    _reorder(wb)
    return ws


def _reorder(wb):
    rank = {n: i for i, n in enumerate(ORDER)}
    wb._sheets.sort(key=lambda ws: rank.get(ws.title, len(ORDER)))
    if "_built" in wb.sheetnames:
        wb["_built"].sheet_state = "hidden"
    # The first visible sheet opens; a hidden one can't be active.
    wb.active = next(i for i, ws in enumerate(wb._sheets) if ws.sheet_state == "visible")
    for ws in wb._sheets:
        ws.sheet_view.tabSelected = ws is wb.active


def style_table(ws, widths, owner_col=None, wrap=()):
    """Header bold and frozen, a filter on every column, widths by header, and
    the Owner column restricted to yes / no with its own colour."""
    head = [c.value for c in ws[1]]
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    ws.freeze_panes = "B2" if len(head) > 6 else "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(head, 1):
        letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[letter].width = widths.get(h, 14)
        if h in wrap:
            for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                row[0].alignment = Alignment(wrap_text=True, vertical="top")
    if owner_col and owner_col in head:
        i = head.index(owner_col) + 1
        letter = openpyxl.utils.get_column_letter(i)
        dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{max(ws.max_row, 2) + 500}")
        for row in ws.iter_rows(min_row=1, min_col=i, max_col=i):
            row[0].fill = OWNER_FILL


def read_sheet(wb, name):
    """(header, [row dicts]) of a sheet, cells normalised to the CSV's text."""
    if name not in wb.sheetnames:
        return [], []
    rows = list(wb[name].iter_rows(values_only=True))
    if not rows:
        return [], []
    head = [norm(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        vals = [norm(v) for v in r]
        if any(vals):
            out.append({h: (vals[i] if i < len(vals) else "") for i, h in enumerate(head) if h})
    return head, out


def save(wb, path=BOOK):
    _reorder(wb)
    wb.save(path)
