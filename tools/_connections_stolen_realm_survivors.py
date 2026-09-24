#!/usr/bin/env python3
"""One-shot: point two `connections` rows at `Stolen Realm Survivors` (connections).

The batch that added Stolen Realm Survivors to `games` also added its two
influences, Diablo and Vampire Survivors, but those rows name the influencee
`Stolen Realm`, which the sheet has no row for. So `import-games-godot.py`
skipped both ("unresolved names: ['Stolen Realm']") and the new game shipped
with nothing pointing at it. Both rows cite the same Reddit post, where the
developers describe a co-op action roguelite they spent two years on. That is
Stolen Realm Survivors (2026), not the 2022 tactics RPG Stolen Realm, and a
Vampire Survivors influence fits only the Survivors game.

Through `replace_cells` for the same reason as `_games_dq_heroes_file_typo.py`:
it touches only the two cells and copies every other zip entry through.

    python3 tools/_connections_stolen_realm_survivors.py
    python3 tools/import-games-godot.py
"""

import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from _xlsx_surgery import Workbook  # noqa: E402

XLSX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Roguelikes.xlsx")

SHEET = "connections"
WAS = "Stolen Realm"
NOW = "Stolen Realm Survivors"
# influencee cell -> the influencer the row must still carry
ROWS = {"B287": ("A287", "Diablo"), "B1267": ("A1267", "Vampire Survivors")}


def main():
    # The refs are only trustworthy if the rows are still the rows this was
    # written against; a sort of the sheet would move them.
    ws = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)[SHEET]
    edits = {}
    for cell, (a_cell, influencer) in ROWS.items():
        a = str(ws[a_cell].value or "").strip()
        b = str(ws[cell].value or "").strip()
        if a != influencer:
            raise SystemExit("%s reads %r, not %r — the sheet has been re-sorted; "
                             "find the row again before rerunning" % (a_cell, a, influencer))
        if b == NOW:
            continue
        if b != WAS:
            raise SystemExit("%s reads %r, not %r — re-read the row before rerunning"
                             % (cell, b, WAS))
        edits[cell] = NOW
    if not edits:
        raise SystemExit("both rows already read %r — nothing to do" % NOW)

    with Workbook(XLSX) as wb:
        wb.replace_cells(SHEET, edits)
    for cell in edits:
        print("connections!%s: %s -> %s" % (cell, WAS, NOW))


if __name__ == "__main__":
    main()
